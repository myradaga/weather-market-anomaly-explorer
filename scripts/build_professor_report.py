from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "Weather_Market_Anomaly_Explorer_Report.docx"
EQ = ROOT / "reports" / "equations"

NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F4F7FB"
GRAY = "D9D9D9"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(89, 99, 114)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = tc_borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr()
    set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        borders(cell)
        margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
        if widths: cell.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if ridx % 2: shade(cell, PALE)
            borders(cell)
            margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths: cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i == 0 and len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for index, item in enumerate(items):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.keep_with_next = index < len(items) - 1
        p.add_run(item)


def add_equation(doc, image_name, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(EQ / image_name), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.style = doc.styles["Caption"]


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.6)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size, before, after in (("Title", 27, 0, 14), ("Heading 1", 17, 13, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11, 8, 4)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = name != "Title"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
styles["Caption"].font.name = "Aptos"
styles["Caption"].font.size = Pt(9)
styles["Caption"].font.italic = True
styles["Caption"].font.color.rgb = MUTED

# Header and footer
header = section.header.paragraphs[0]
header.text = "Weather Market Anomaly Explorer   Research Methods Report"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8.5)
header.runs[0].font.color.rgb = MUTED
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Page ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)

# Cover page
doc.add_paragraph("Weather Market Trading Anomaly Explorer", style="Title")
subtitle = doc.add_paragraph("Research methods and system report")
subtitle.style = styles["Subtitle"]
subtitle.runs[0].font.color.rgb = MUTED
doc.add_paragraph("Polymarket and Kalshi weather prediction markets")
doc.add_paragraph("Prepared for academic review\n1 October 2026")
doc.add_paragraph().add_run("Purpose").bold = True
doc.add_paragraph(
    "This report explains what the application does, how it calculates each score, and how to read the results. The system finds unusual public trading activity that may deserve human review. It does not decide that fraud or illegal trading occurred."
)
doc.add_paragraph().add_run("Main point").bold = True
doc.add_paragraph(
    "The application turns a large number of public trades into a shorter list of cases to review. Each case shows the trades, calculations, signal scores, and missing information. Unusual activity is not proof of wrongdoing. Kalshi also does not provide public trader identities."
)
add_table(doc, ["Platform", "Unit of analysis", "Identity evidence", "Appropriate conclusion"], [
    ["Polymarket", "Wallet market episode", "Pseudonymous public wallet", "Which wallet episodes deserve review"],
    ["Kalshi", "Trade and grouped market activity event", "No public trader identity", "Which market activity periods deserve review"],
], [1.15, 1.8, 1.7, 2.25], 9.2)
doc.add_page_break()

doc.add_heading("Purpose of the project", level=1)
doc.add_paragraph(
    "The project asks a simple question: which trades in open weather prediction markets look unusual enough for a researcher to review? It collects public market data, groups related trades, calculates five signals, and shows the strongest cases in a dashboard."
)
doc.add_paragraph(
    "A higher score means that several parts of the trading pattern are unusual compared with similar activity. The score is not the probability of fraud. A researcher must still check market rules, public news or weather updates, data quality, and other possible explanations."
)

doc.add_heading("System workflow", level=1)
add_table(doc, ["Stage", "What the system does", "Research value"], [
    ["1", "Finds open weather markets and removes nonweather markets", "Keeps the project focused"],
    ["2", "Downloads public markets, prices, and trades", "Creates a record of the source data"],
    ["3", "Standardizes times, outcomes, prices, sizes, and dollar values", "Makes trades easier to compare"],
    ["4", "Groups related activity into Polymarket episodes and Kalshi events", "Stops each small trade from becoming its own case"],
    ["5", "Calculates five signals and checks data reliability", "Shows why an activity pattern looks unusual"],
    ["6", "Uses strict rules before showing a flag", "Reduces weak or misleading flags"],
    ["7", "Shows graphs, calculations, limits, and source links", "Helps a researcher review the evidence"],
], [0.55, 3.25, 2.9], 9.1)

doc.add_heading("Important terminology", level=1)
add_table(doc, ["Term", "Meaning in this project"], [
    ["Market", "A question people can trade on; its price shows the market's current view of the outcome"],
    ["Wallet", "A public pseudonymous blockchain address on Polymarket; it is not a verified personal identity"],
    ["Trade notional", "The dollar value of an executed trade, calculated as price multiplied by contracts or shares"],
    ["Episode", "One wallet's activity in one Polymarket market, separated from later activity by a 24 hour inactivity gap"],
    ["Kalshi activity event", "Related trades in one contract grouped until a 30 minute inactivity gap"],
    ["Signal score", "A value from 0 to 1; a higher value means that part of the activity is more unusual"],
    ["Reliability weight", "A value that lowers or removes a signal when the supporting data are weak or missing"],
], [1.55, 5.15], 9.25)
add_equation(doc, "notional.png", "Equation 1  Trade notional", 4.4)

doc.add_heading("Polymarket wallet episode model", level=1)
doc.add_paragraph(
    "Polymarket includes public wallet addresses with its trade data. The application can group one wallet's trades in one market and compare that episode with other traders and with the wallet's earlier activity. The five signals are called S1 through S5."
)
add_table(doc, ["Signal", "Question", "Calculation and interpretation"], [
    ["S1 Size", "Is this episode unusually large", "Compares the dollar amount with other episodes in the market and with the wallet's earlier episodes. The larger comparison score is used."],
    ["S2 Timing", "Did the trade happen at an unusual time", "Compares the trade time with a verified public weather update when one is available. A weaker close-time estimate is used when verified timing is missing."],
    ["S3 Past results", "Has the wallet done unusually well before", "Uses earlier resolved results. The score stays close to neutral when the wallet has little history."],
    ["S4 Market share", "Did the wallet make up a large share of activity", "Combines the market's YES versus NO imbalance with the wallet's share of observed trading. Very small markets are excluded."],
    ["S5 Behavior", "Does the overall pattern differ from similar episodes", "Uses an Isolation Forest model to compare size, market share, direction, and earlier activity with similar weather episodes."],
], [0.9, 1.75, 4.05], 8.75)
add_equation(doc, "poly_signals.png", "Equation 2  Selected Polymarket signal formulas", 6.4)

doc.add_heading("Reliability controls", level=2)
add_bullets(doc, [
    "S2 receives full weight only when a trusted public timestamp supports the timing comparison.",
    "S3 stays closer to neutral when only a few earlier wallet results are available.",
    "S4 is not used when the market has fewer than 20 trades or less than 1000 dollars in observed trading.",
    "A missing signal is left out instead of being scored as zero. Zero would incorrectly suggest normal activity.",
])

doc.add_heading("Composite anomaly score", level=1)
doc.add_paragraph(
    "The system does not simply average the five signals. It changes each available signal into log odds, gives weaker data less weight, and combines the results. The logistic function then changes the result back to a score between 0 and 1. Several strong signals together raise the score, while weaker or opposite signals lower it."
)
add_equation(doc, "composite.png", "Equation 3  Reliability weighted evidence combination", 6.5)
doc.add_paragraph(
    "In the formula, Si is one signal score, Qi is the reliability of that signal, and sigma changes the result back to the 0 to 1 scale. This is an anomaly score, not a probability of fraud."
)

doc.add_heading("Polymarket review thresholds", level=2)
add_table(doc, ["Requirement", "Implemented rule"], [
    ["Composite score", "At least 0.75"],
    ["Valid evidence", "At least four available modules"],
    ["Strong evidence", "At least three active signals at or above 0.75 in the current scoring code"],
    ["Minimum activity", "At least two trades and 100 dollars notional, or one episode of at least 5000 dollars"],
    ["Priority band", "Score at least 0.92, all five modules available, and high confidence timing evidence"],
], [2.0, 4.7], 9.2)
doc.add_page_break()

doc.add_heading("Kalshi market activity model", level=1)
doc.add_paragraph(
    "Kalshi's public data does not include a trader or wallet identifier. The application therefore cannot identify a person or study one account's history. Instead, it looks only at visible market activity. Kalshi flags describe unusual trades or market events, not suspicious users."
)
add_table(doc, ["Signal", "Public evidence used", "Main limitation"], [
    ["K1 Trade size", "Compares the trade with similar trades and checks whether it is at least meaningful in dollars", "A trade can be large compared with peers but still small in dollars"],
    ["K2 Timing", "Looks for a burst of trades and activity close to market closing", "It receives less weight without a verified public weather-update time"],
    ["K3 Price movement", "Looks for a price jump followed by a move back", "A price move does not show the trader's reason"],
    ["K4 Market pressure", "Measures the trade as a share of observed activity in the contract", "Public data may not show every part of market liquidity"],
    ["K5 Difference from peers", "Compares size, price, contract count, and time gaps with similar weather markets", "An unusual pattern does not explain its cause"],
], [1.15, 3.25, 2.3], 8.7)
add_equation(doc, "kalshi_signals.png", "Equation 4  Selected Kalshi signal formulas", 6.4)

doc.add_heading("Kalshi queue and event grouping", level=2)
doc.add_paragraph(
    "A Kalshi trade is shown for review only when its total score is at least 0.95, its size score is at least 0.85, at least three signals are 0.80 or higher, and the trade value is at least 250 dollars. Related trades are grouped into one event when they occur less than 30 minutes apart."
)
doc.add_paragraph(
    "A high evidence grade requires at least 1000 dollars in grouped trading, four strong signals, and a meaningful price movement. A medium grade requires at least 250 dollars and three strong signals. These grades help order the review list. They do not prove misconduct."
)

doc.add_heading("Evidence unavailable for Kalshi", level=2)
add_table(doc, ["Unavailable evidence", "What cannot be concluded"], [
    ["Trader identity", "The dashboard cannot identify a person or account"],
    ["Account history", "It cannot calculate repeat success by the same user"],
    ["Counterparty links", "It cannot detect coordinated accounts"],
    ["Order submissions and cancellations", "It cannot infer intent from unexecuted orders"],
    ["Account ownership on both sides", "It cannot perform a direct wash trading test"],
    ["Full order book behavior", "It cannot establish spoofing or layering"],
], [2.35, 4.35], 9.1)

doc.add_heading("How to read a flagged case", level=1)
add_table(doc, ["Dashboard element", "What to ask"], [
    ["Market question", "What exact outcome is being priced, and what are the settlement rules"],
    ["Progression graph", "Did traded value cluster in time, and did the executed price move or reverse"],
    ["Trade table", "Does price multiplied by size reproduce the displayed notional"],
    ["Signal table", "Which signals are high, and which reliability weights weaken them"],
    ["Timing reference", "Is the public event timestamp verified, or is it only a close time proxy"],
    ["Missing evidence", "Which conclusions are impossible with the available public data"],
    ["Cross market context", "Do similar contracts genuinely share location, unit, time period, and settlement rules"],
], [2.05, 4.65], 9.2)

doc.add_heading("Worked interpretation example", level=1)
doc.add_paragraph(
    "Suppose a Kalshi event contains 2500 dollars across several trades. Four of the five signals are above 0.80, and the price rises from 0.30 to 0.75 before falling to 0.45. The application may give this event a high evidence grade. This means the burst was large, the price move was unusual, and the event differed from similar activity. It does not tell us who traded, why they traded, or whether the activity was manipulative."
)

doc.add_heading("Research safeguards and limitations", level=1)
add_bullets(doc, [
    "The system uses public data and does not try to discover the real people behind Polymarket wallets.",
    "Open markets can change after data is collected, so each case should include a collection time.",
    "A market price shows traders' current views. It is not a direct measurement of the true probability.",
    "Isolation Forest finds rare patterns. It does not explain why the pattern is rare.",
    "Public weather-update times must be verified. A market's close time alone is not proof of an information advantage.",
    "The current thresholds are research rules. They have not been tested against a confirmed fraud dataset.",
    "A comparison across platforms is only useful when the markets have similar rules, units, locations, and time periods.",
])

doc.add_heading("Validation and reproducibility", level=1)
doc.add_paragraph(
    "The project saves the original data, uses stable record IDs, and keeps data collection separate from scoring. It also records the feature and model versions. The automated test suite has 18 passing tests covering data clients, trade grouping, calculations, scoring, Kalshi events, and basic application behavior. These tests show that the software runs consistently, but they do not prove that every flag is correct."
)

doc.add_heading("Improvements from the earlier version", level=1)
add_table(doc, ["Earlier version", "Improvement made"], [
    ["Included many market categories", "Now focuses only on weather markets and removes nonweather Kalshi results"],
    ["Mixed open and resolved markets", "Main research screens now focus on open or live markets"],
    ["Used demo or synthetic records", "Dashboard now uses live public Polymarket and Kalshi data instead of demo data"],
    ["Polymarket and Kalshi were presented together", "Each platform now has its own section, method, formulas, and limits"],
    ["Flags could appear with weak visible evidence", "Stricter score, signal-count, trade-size, and data-quality rules now control the review list"],
    ["Kalshi lacked wallet information", "A separate market-activity model now scores only the public evidence Kalshi provides"],
    ["A flag was difficult to understand", "Each case now shows why it was flagged, the trade progression graph, calculations, reliability, and missing evidence"],
    ["The interface was crowded and hard to read", "The layout now uses clearer labels, score keys, readable graph text, separate platform views, and smaller status indicators"],
], [2.25, 4.45], 8.9)

doc.add_heading("Possible next improvements", level=1)
doc.add_paragraph(
    "Two changes would make the project stronger. First, verified weather-update times from sources such as NOAA or the National Weather Service would improve the timing signal. Second, a set of cases reviewed by people would help measure false positives and choose better score thresholds."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
